"""Business-analysis PPTX deck, built with native chart objects
(CategoryChartData -> real, editable PowerPoint charts, not embedded chart
screenshots) from data already computed by domain/analytics.py and
domain/inventory.py.

The insight bullets on the closing slide are computed here, from the data
-- never phrased or invented by the model. The model may relay them
conversationally in chat, but the deck's own text is always arithmetic on
real numbers, so it can't be wrong in the way an LLM-authored summary could."""

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "data" / "output"

_TITLE_LAYOUT = 0
_BLANK_LAYOUT = 6


def _add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[_TITLE_LAYOUT])
    slide.shapes.title.text = title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle
    return slide


def _add_blank_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT])
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(9.2), Inches(0.7))
    tf = box.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.bold = True
    return slide


def _empty_notice(slide, text):
    slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(8), Inches(1)).text_frame.text = text


def _kpi_slide(prs, summary):
    slide = _add_blank_slide(prs, "Key Numbers")
    avg_bill = summary["total_sales"] / summary["bill_count"] if summary["bill_count"] else 0
    rows = [
        ("Total Sales", f"Rs. {summary['total_sales']:,.2f}"),
        ("Bills", str(summary["bill_count"])),
        ("Tax Collected", f"Rs. {summary['tax_collected']:,.2f}"),
        ("Avg. Bill Value", f"Rs. {avg_bill:,.2f}"),
    ]
    table = slide.shapes.add_table(len(rows), 2, Inches(0.6), Inches(1.3), Inches(8.5), Inches(2.2)).table
    for i, (label, value) in enumerate(rows):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value
    return slide


def _top_items_slide(prs, top_items):
    slide = _add_blank_slide(prs, "Top Items by Revenue")
    if not top_items:
        _empty_notice(slide, "No sales in this period.")
        return slide
    chart_data = CategoryChartData()
    chart_data.categories = [item["name"] for item in top_items]
    chart_data.add_series("Revenue (Rs.)", [round(item["revenue"] or 0, 2) for item in top_items])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.2), Inches(8.5), Inches(4.8), chart_data)
    return slide


def _daily_trend_slide(prs, daily_trend):
    slide = _add_blank_slide(prs, "Daily Sales Trend")
    if not daily_trend:
        _empty_notice(slide, "No sales in this period.")
        return slide
    chart_data = CategoryChartData()
    chart_data.categories = [row["d"] for row in daily_trend]
    chart_data.add_series("Sales (Rs.)", [round(row["total"] or 0, 2) for row in daily_trend])
    slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.6), Inches(1.2), Inches(8.5), Inches(4.8), chart_data)
    return slide


def _payment_mode_slide(prs, payment_mode_split):
    slide = _add_blank_slide(prs, "Payment Mode Split")
    if not payment_mode_split:
        _empty_notice(slide, "No sales in this period.")
        return slide
    chart_data = CategoryChartData()
    chart_data.categories = [k.upper() for k in payment_mode_split.keys()]
    chart_data.add_series("Share", list(payment_mode_split.values()))
    graphic_frame = slide.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(1.5), Inches(1.2), Inches(6), Inches(4.8), chart_data)
    chart = graphic_frame.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    return slide


def _stock_health_slide(prs, reorder_suggestions):
    slide = _add_blank_slide(prs, "Stock Health / Reorder Suggestions")
    if not reorder_suggestions:
        _empty_notice(slide, "Nothing needs reordering right now.")
        return slide
    rows = [("Item", "On Hand", "Reorder Level", "Est. Days Left")] + [
        (
            s["name"], f"{s['qty_on_hand']:g} {s['unit']}", f"{s['reorder_level']:g}",
            str(s["estimated_days_left"]) if s["estimated_days_left"] is not None else "-",
        )
        for s in reorder_suggestions
    ]
    table = slide.shapes.add_table(
        len(rows), 4, Inches(0.5), Inches(1.2), Inches(9), Inches(min(0.4 * len(rows), 5))
    ).table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = str(val)
    return slide


def _build_insights(summary: dict, reorder_suggestions: list[dict]) -> list[str]:
    """Pure arithmetic over already-computed data -- no model call."""
    if not summary["bill_count"]:
        return ["No sales recorded in this period."]

    insights = [
        f"{summary['bill_count']} bill(s) totalling Rs. {summary['total_sales']:,.2f}, "
        f"with Rs. {summary['tax_collected']:,.2f} in GST collected."
    ]
    if summary["top_items"]:
        top = summary["top_items"][0]
        insights.append(f"Top seller by revenue: {top['name']} (Rs. {(top['revenue'] or 0):,.2f}).")

    split = summary["payment_mode_split"]
    if split:
        total = sum(split.values()) or 1
        dominant_mode, dominant_amt = max(split.items(), key=lambda kv: kv[1])
        insights.append(f"{dominant_mode.upper()} led payment mode at {dominant_amt / total * 100:.0f}% of sales.")

    if reorder_suggestions:
        names = ", ".join(s["name"] for s in reorder_suggestions[:5])
        insights.append(f"{len(reorder_suggestions)} item(s) need reordering soon: {names}.")
    else:
        insights.append("No items currently need reordering.")
    return insights


def _insights_slide(prs, insights: list[str]):
    slide = _add_blank_slide(prs, "Insights")
    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(8.5), Inches(5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(insights):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {line}"
        p.font.size = Pt(18)
    return slide


def render_analysis_deck(
    shop_name: str, start_date: str, end_date: str, summary: dict, reorder_suggestions: list[dict]
) -> str:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / f"analysis_{start_date}_to_{end_date}.pptx"

    prs = Presentation()
    _add_title_slide(prs, f"{shop_name} -- Sales Analysis", f"{start_date} to {end_date}")
    _kpi_slide(prs, summary)
    _top_items_slide(prs, summary["top_items"])
    _daily_trend_slide(prs, summary["daily_trend"])
    _payment_mode_slide(prs, summary["payment_mode_split"])
    _stock_health_slide(prs, reorder_suggestions)
    _insights_slide(prs, _build_insights(summary, reorder_suggestions))

    prs.save(str(path))
    return str(path)
